import os
import re
import PyPDF2
import docx
import csv
import openpyxl

def extract_emails(text):
    email_matches = re.findall(r'\S+@\S+', text)
    return email_matches

def extract_phone_numbers(text):
    phone_matches = re.findall(r'\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b', text)
    return phone_matches

def extract_web_urls(text):
    url_matches = re.findall(r'http[s]?://(?:[a-zA-Z]|[0-9]|[$-_@.&+]|[!*\\(\\),]|(?:%[0-9a-fA-F][0-9a-fA-F]))+', text)
    return url_matches

def process_txt(txt_path):
    with open(txt_path, 'r') as txt_file:
        text = txt_file.read()
        extracted_info = {
            "emails": extract_emails(text),
            "phone_numbers": extract_phone_numbers(text),
            "web_urls": extract_web_urls(text)
        }
        return extracted_info

def process_pdf(pdf_path):
    extracted_info = {
        "emails": [],
        "phone_numbers": [],
        "web_urls": []
    }

    try:
        pdf_file = open(pdf_path, "rb")
        pdf_reader = PyPDF2.PdfReader(pdf_file)

        for page_num in range(len(pdf_reader.pages)):
            page = pdf_reader.pages[page_num]
            text = page.extract_text()

            extracted_info["emails"].extend(extract_emails(text))
            extracted_info["phone_numbers"].extend(extract_phone_numbers(text))
            extracted_info["web_urls"].extend(extract_web_urls(text))

        pdf_file.close()
    except Exception as e:
        print("An error occurred:", e)

    return extracted_info

def process_docx(docx_path):
    extracted_info = {
        "emails": [],
        "phone_numbers": [],
        "web_urls": []
    }

    try:
        doc = docx.Document(docx_path)

        for paragraph in doc.paragraphs:
            text = paragraph.text

            extracted_info["emails"].extend(extract_emails(text))
            extracted_info["phone_numbers"].extend(extract_phone_numbers(text))
            extracted_info["web_urls"].extend(extract_web_urls(text))

    except Exception as e:
        print("An error occurred:", e)

    return extracted_info

def process_csv(csv_path):
    extracted_info = {
        "emails": [],
        "phone_numbers": [],
        "web_urls": []
    }

    try:
        with open(csv_path, 'r') as csv_file:
            csv_reader = csv.reader(csv_file)
            for row in csv_reader:
                text = ' '.join(row)

                extracted_info["emails"].extend(extract_emails(text))
                extracted_info["phone_numbers"].extend(extract_phone_numbers(text))
                extracted_info["web_urls"].extend(extract_web_urls(text))
    except Exception as e:
        print("An error occurred:", e)

    return extracted_info

def process_excel(excel_path):
    extracted_info = {
        "emails": [],
        "phone_numbers": [],
        "web_urls": []
    }

    try:
        wb = openpyxl.load_workbook(excel_path)
        for sheet in wb:
            for row in sheet.iter_rows():
                for cell in row:
                    text = str(cell.value)

                    extracted_info["emails"].extend(extract_emails(text))
                    extracted_info["phone_numbers"].extend(extract_phone_numbers(text))
                    extracted_info["web_urls"].extend(extract_web_urls(text))
    except Exception as e:
        print("An error occurred:", e)

    return extracted_info

def main():
    folder_path = input("Enter the folder path: ")
    ext_mapping = {
        'txt': process_txt,
        'pdf': process_pdf,
        'docx': process_docx,
        'csv': process_csv,
        'xlsx': process_excel
        # Add more mappings as needed for other formats
    }

    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        ext = filename.split('.')[-1].lower()

        if ext in ext_mapping:
            process_func = ext_mapping[ext]
            extracted_info = process_func(file_path)

            print(f"Extracted info from {filename}:")
            print("Emails:", extracted_info["emails"])
            print("Phone Numbers:", extracted_info["phone_numbers"])
            print("Web URLs:", extracted_info["web_urls"])
            print("=" * 50)

if __name__ == "__main__":
    main()
